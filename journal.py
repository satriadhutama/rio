"""
Quant Edge Engine — Crypto Trade Journal
Baca & tulis ke quant_edge_journal.xlsx
"""

import sys
import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

EXCEL_FILE = "quant_edge_journal.xlsx"

# ── Konstanta validasi ──────────────────────────────────────────────────────
VALID_SETUP     = {"Narrative", "Divergence", "SM pre-pump", "Lain"}
VALID_BIAS      = {"long", "short"}
VALID_CONFIDENCE= {"Low", "Med", "High"}
VALID_VERDICT   = {"EKSEKUSI", "WAIT", "ABAIKAN"}
VALID_HASIL     = {"Win", "Loss", "BE", "Skip", "Pending"}
VALID_PILAR     = {-1, 0, 1}

# ── Header sheet Log ────────────────────────────────────────────────────────
LOG_HEADERS = [
    "No", "Tanggal", "Coin", "Setup", "Bias", "Score", "Confidence",
    "Verdict", "Entry", "SL", "TP1", "TP2", "TP3", "RR_plan",
    "P1", "P2", "P3", "P4", "P5", "P6",
    "Diambil", "Hasil", "Exit", "R_aktual", "Invalidasi_kena", "Catatan"
]

# ── Warna ──────────────────────────────────────────────────────────────────
GREEN_FILL  = PatternFill("solid", fgColor="C6EFCE")
RED_FILL    = PatternFill("solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
HEADER_FILL = PatternFill("solid", fgColor="2F4F8F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BOLD_FONT   = Font(bold=True)

thin = Side(style="thin")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# ═══════════════════════════════════════════════════════════════════════════
#  INISIALISASI / BUKA FILE
# ═══════════════════════════════════════════════════════════════════════════

def open_or_create_workbook():
    """Buka Excel jika ada; buat baru jika belum ada."""
    if os.path.exists(EXCEL_FILE):
        try:
            wb = load_workbook(EXCEL_FILE)
        except PermissionError:
            print(f"\n[ERROR] File '{EXCEL_FILE}' sedang dibuka di Excel.")
            print("        Tutup dulu file-nya, lalu jalankan ulang.\n")
            sys.exit(1)
        # Pastikan ketiga sheet ada
        for name in ("Log", "Evaluasi", "Equity"):
            if name not in wb.sheetnames:
                wb.create_sheet(name)
        return wb
    else:
        wb = Workbook()
        # Hapus sheet default
        for s in wb.sheetnames:
            del wb[s]
        _init_log_sheet(wb.create_sheet("Log"))
        wb.create_sheet("Evaluasi")
        wb.create_sheet("Equity")
        return wb


def _init_log_sheet(ws):
    """Tulis header ke sheet Log dengan styling."""
    ws.append(LOG_HEADERS)
    for col, hdr in enumerate(LOG_HEADERS, 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BORDER
    # Lebar kolom
    widths = {
        "No":2, "Tanggal":12, "Coin":6, "Setup":14, "Bias":5,
        "Score":5, "Confidence":10, "Verdict":10, "Entry":7, "SL":7,
        "TP1":7, "TP2":7, "TP3":7, "RR_plan":7,
        "P1":3,"P2":3,"P3":3,"P4":3,"P5":3,"P6":3,
        "Diambil":7, "Hasil":8, "Exit":7, "R_aktual":8,
        "Invalidasi_kena":13, "Catatan":30,
    }
    for i, hdr in enumerate(LOG_HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(hdr, 10)
    ws.freeze_panes = "A2"


# ═══════════════════════════════════════════════════════════════════════════
#  BACA DATA LOG
# ═══════════════════════════════════════════════════════════════════════════

def read_log(wb):
    """Kembalikan list of dict dari sheet Log (skip baris header)."""
    ws = wb["Log"]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        rows.append(dict(zip(LOG_HEADERS, row)))
    return rows


def next_no(wb):
    rows = read_log(wb)
    if not rows:
        return 1
    return max(int(r["No"]) for r in rows) + 1


# ═══════════════════════════════════════════════════════════════════════════
#  VALIDASI
# ═══════════════════════════════════════════════════════════════════════════

def validate_add(d):
    errors = []
    if d.get("Setup") and d["Setup"] not in VALID_SETUP:
        errors.append(f"  Setup '{d['Setup']}' tidak valid. Pilihan: {sorted(VALID_SETUP)}")
    if d.get("Bias") and d["Bias"] not in VALID_BIAS:
        errors.append(f"  Bias '{d['Bias']}' tidak valid. Pilihan: {sorted(VALID_BIAS)}")
    if d.get("Confidence") and d["Confidence"] not in VALID_CONFIDENCE:
        errors.append(f"  Confidence '{d['Confidence']}' tidak valid. Pilihan: {sorted(VALID_CONFIDENCE)}")
    if d.get("Verdict") and d["Verdict"] not in VALID_VERDICT:
        errors.append(f"  Verdict '{d['Verdict']}' tidak valid. Pilihan: {sorted(VALID_VERDICT)}")
    for p in ("P1","P2","P3","P4","P5","P6"):
        v = d.get(p)
        if v is not None and v != "" and int(v) not in VALID_PILAR:
            errors.append(f"  {p}='{v}' tidak valid. Nilai harus -1, 0, atau 1.")
    return errors


def validate_hasil(d):
    errors = []
    if d.get("Hasil") and d["Hasil"] not in VALID_HASIL:
        errors.append(f"  Hasil '{d['Hasil']}' tidak valid. Pilihan: {sorted(VALID_HASIL)}")
    return errors


# ═══════════════════════════════════════════════════════════════════════════
#  TULIS BARIS BARU KE LOG
# ═══════════════════════════════════════════════════════════════════════════

def append_trade(wb, d):
    ws = wb["Log"]
    if ws.max_row == 1 and ws.cell(1,1).value != "No":
        _init_log_sheet(ws)

    row_data = [d.get(h, "") for h in LOG_HEADERS]
    ws.append(row_data)

    # Styling baris baru
    row_num = ws.max_row
    hasil = d.get("Hasil", "Pending")
    if hasil == "Win":
        fill = GREEN_FILL
    elif hasil == "Loss":
        fill = RED_FILL
    elif hasil == "BE":
        fill = YELLOW_FILL
    else:
        fill = None

    for col in range(1, len(LOG_HEADERS)+1):
        cell = ws.cell(row_num, col)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
        if fill:
            cell.fill = fill


def update_trade(wb, no, update_dict):
    """Update baris di sheet Log berdasarkan No trade."""
    ws = wb["Log"]
    header_row = {cell.value: cell.column for cell in ws[1]}

    for row in ws.iter_rows(min_row=2):
        if row[0].value == no:
            for key, val in update_dict.items():
                if key in header_row:
                    ws.cell(row[0].row, header_row[key]).value = val
            # Warna ulang baris
            hasil = ws.cell(row[0].row, header_row["Hasil"]).value
            if hasil == "Win":
                fill = GREEN_FILL
            elif hasil == "Loss":
                fill = RED_FILL
            elif hasil == "BE":
                fill = YELLOW_FILL
            else:
                fill = None
            if fill:
                for col in range(1, len(LOG_HEADERS)+1):
                    ws.cell(row[0].row, col).fill = fill
            return True
    return False


# ═══════════════════════════════════════════════════════════════════════════
#  KALKULASI METRIK
# ═══════════════════════════════════════════════════════════════════════════

def _safe_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _safe_int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def compute_metrics(rows):
    """Hitung semua metrik dari list of dict. Return dict metrik."""
    finished = [r for r in rows if r.get("Hasil") not in (None, "", "Pending", "Skip")]
    pending  = [r for r in rows if r.get("Hasil") in (None, "", "Pending")]
    skipped  = [r for r in rows if r.get("Hasil") == "Skip"]

    wins   = [r for r in finished if r.get("Hasil") == "Win"]
    losses = [r for r in finished if r.get("Hasil") == "Loss"]
    bes    = [r for r in finished if r.get("Hasil") == "BE"]

    n_fin  = len(finished)
    win_rate = len(wins) / n_fin if n_fin else 0
    loss_rate= (len(losses)+len(bes)) / n_fin if n_fin else 0  # BE = 0R, masih loss-side

    r_wins   = [_safe_float(r["R_aktual"]) for r in wins if _safe_float(r["R_aktual"]) is not None]
    r_losses = [abs(_safe_float(r["R_aktual"])) for r in losses if _safe_float(r["R_aktual"]) is not None]

    avg_win  = sum(r_wins) / len(r_wins) if r_wins else 0
    avg_loss = sum(r_losses) / len(r_losses) if r_losses else 0

    # Expectancy = WR*AvgWin - LossRate*AvgLoss  (BE di-treat sebagai 0R)
    be_loss_rate = len(losses) / n_fin if n_fin else 0
    expectancy = (win_rate * avg_win) - (be_loss_rate * avg_loss) if n_fin else 0

    all_r = [_safe_float(r["R_aktual"]) for r in finished if _safe_float(r["R_aktual"]) is not None]
    total_r = sum(all_r)

    # ── Per Pilar ──────────────────────────────────────────────────────
    pilar_stats = {}
    for p in ("P1","P2","P3","P4","P5","P6"):
        active = [r for r in finished if _safe_int(r.get(p)) == 1]
        w = [r for r in active if r.get("Hasil") == "Win"]
        n = len(active)
        wr = len(w)/n if n else None
        pilar_stats[p] = {"wr": wr, "n": n}

    # ── Per Setup ──────────────────────────────────────────────────────
    setup_stats = {}
    for setup in VALID_SETUP:
        grp = [r for r in finished if r.get("Setup") == setup]
        if not grp:
            continue
        w = [r for r in grp if r.get("Hasil") == "Win"]
        l = [r for r in grp if r.get("Hasil") == "Loss"]
        n = len(grp)
        wr = len(w)/n
        lr = len(l)/n
        rw = [_safe_float(r["R_aktual"]) for r in w if _safe_float(r["R_aktual"]) is not None]
        rl = [abs(_safe_float(r["R_aktual"])) for r in l if _safe_float(r["R_aktual"]) is not None]
        aw = sum(rw)/len(rw) if rw else 0
        al = sum(rl)/len(rl) if rl else 0
        exp = wr*aw - lr*al
        setup_stats[setup] = {"wr": wr, "n": n, "exp": exp}

    # ── Per Verdict ────────────────────────────────────────────────────
    verdict_stats = {}
    for v in VALID_VERDICT:
        grp = [r for r in finished if r.get("Verdict") == v]
        if not grp:
            continue
        w = [r for r in grp if r.get("Hasil") == "Win"]
        l = [r for r in grp if r.get("Hasil") == "Loss"]
        n = len(grp)
        wr = len(w)/n
        lr = len(l)/n
        rw = [_safe_float(r["R_aktual"]) for r in w if _safe_float(r["R_aktual"]) is not None]
        rl = [abs(_safe_float(r["R_aktual"])) for r in l if _safe_float(r["R_aktual"]) is not None]
        aw = sum(rw)/len(rw) if rw else 0
        al = sum(rl)/len(rl) if rl else 0
        exp = wr*aw - lr*al
        verdict_stats[v] = {"wr": wr, "n": n, "exp": exp}

    # ── Equity curve ──────────────────────────────────────────────────
    equity = []
    cum = 0.0
    for r in finished:
        rv = _safe_float(r.get("R_aktual"))
        if rv is not None:
            cum += rv
        equity.append(round(cum, 2))

    return {
        "total":       len(rows),
        "finished":    n_fin,
        "pending":     len(pending),
        "skipped":     len(skipped),
        "wins":        len(wins),
        "losses":      len(losses),
        "bes":         len(bes),
        "win_rate":    win_rate,
        "avg_win":     avg_win,
        "avg_loss":    avg_loss,
        "expectancy":  expectancy,
        "total_r":     total_r,
        "pilar":       pilar_stats,
        "setup":       setup_stats,
        "verdict":     verdict_stats,
        "equity":      equity,
        "finished_rows": finished,
    }


# ═══════════════════════════════════════════════════════════════════════════
#  TULIS SHEET EVALUASI
# ═══════════════════════════════════════════════════════════════════════════

def _cell(ws, row, col, value, bold=False, fill=None, fmt=None, align="left"):
    c = ws.cell(row, col, value)
    c.border = BORDER
    c.alignment = Alignment(horizontal=align, wrap_text=True)
    if bold:
        c.font = BOLD_FONT
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def write_evaluasi(wb, m):
    ws = wb["Evaluasi"]
    ws.delete_rows(1, ws.max_row + 1)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 16

    r = 1

    # ── Judul ──
    ws.merge_cells(f"A{r}:E{r}")
    c = ws.cell(r, 1, "QUANT EDGE ENGINE — EVALUASI JOURNAL")
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center")
    r += 1

    ws.cell(r, 1, f"Diupdate: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = Font(italic=True, color="888888")
    r += 2

    # ── EXPECTANCY (paling atas, beri warna) ──
    exp_fill = GREEN_FILL if m["expectancy"] >= 0 else RED_FILL
    _cell(ws, r, 1, "★ EXPECTANCY per Trade (R)", bold=True, fill=exp_fill)
    _cell(ws, r, 2, round(m["expectancy"], 2), bold=True, fill=exp_fill, fmt="0.00", align="center")
    if m["expectancy"] >= 0:
        _cell(ws, r, 3, "✔ Positif — sistem layak dioperasikan", fill=exp_fill)
    else:
        _cell(ws, r, 3, "✘ Negatif — perlu review sistem", fill=exp_fill)
    r += 2

    # ── Ringkasan umum ──
    def kv(label, val, row, fmt=None):
        _cell(ws, row, 1, label, bold=True)
        _cell(ws, row, 2, val, fmt=fmt, align="center")

    kv("Total Trade (semua)", m["total"],    r); r+=1
    kv("Trade Selesai",        m["finished"], r); r+=1
    kv("Pending",              m["pending"],  r); r+=1
    kv("Skipped",              m["skipped"],  r); r+=1
    kv("Win",                  m["wins"],     r); r+=1
    kv("Loss",                 m["losses"],   r); r+=1
    kv("Break Even",           m["bes"],      r); r+=2

    kv("Win Rate",    f"{m['win_rate']*100:.1f}%",  r); r+=1
    kv("Avg Win (R)", round(m["avg_win"],2),         r, "0.00"); r+=1
    kv("Avg Loss (R)",round(m["avg_loss"],2),        r, "0.00"); r+=1
    kv("Total R Kumulatif", round(m["total_r"],2),   r, "0.00"); r+=2

    # ── Per Pilar ──
    ws.merge_cells(f"A{r}:E{r}")
    ch = ws.cell(r,1,"ANALISIS PER PILAR (P1–P6)")
    ch.font = Font(bold=True, color="FFFFFF"); ch.fill = HEADER_FILL
    ch.alignment = Alignment(horizontal="center")
    r+=1

    for hdr, col in [("Pilar",1),("Win Rate",2),("Sampel (n)",3),("Keterangan",4)]:
        _cell(ws, r, col, hdr, bold=True, fill=PatternFill("solid",fgColor="D9E1F2"))
    r+=1

    for p in ("P1","P2","P3","P4","P5","P6"):
        st = m["pilar"][p]
        wr = st["wr"]
        n  = st["n"]
        if wr is None:
            wr_str = "—"
            ket    = "Belum ada data"
            fill   = None
        else:
            wr_str = f"{wr*100:.1f}%"
            if n < 10:
                ket  = "(sampel kecil, belum reliabel)"
                fill = YELLOW_FILL
            elif wr >= 0.55:
                ket  = "Prediktif — bobot layak naik"
                fill = GREEN_FILL
            elif wr < 0.45:
                ket  = "Lemah — bobot layak turun"
                fill = RED_FILL
            else:
                ket  = "Netral"
                fill = None
        _cell(ws, r, 1, p,      bold=True, align="center")
        _cell(ws, r, 2, wr_str, align="center", fill=fill)
        _cell(ws, r, 3, n,      align="center")
        _cell(ws, r, 4, ket,    fill=fill)
        r+=1
    r+=1

    # ── Per Setup ──
    ws.merge_cells(f"A{r}:E{r}")
    ch = ws.cell(r,1,"ANALISIS PER SETUP")
    ch.font = Font(bold=True,color="FFFFFF"); ch.fill=HEADER_FILL
    ch.alignment = Alignment(horizontal="center")
    r+=1

    for hdr,col in [("Setup",1),("Win Rate",2),("Expectancy (R)",3),("Sampel (n)",4)]:
        _cell(ws,r,col,hdr,bold=True,fill=PatternFill("solid",fgColor="D9E1F2"))
    r+=1

    for setup, st in m["setup"].items():
        exp_fill2 = GREEN_FILL if st["exp"]>=0 else RED_FILL
        _cell(ws,r,1,setup)
        _cell(ws,r,2,f"{st['wr']*100:.1f}%", align="center")
        _cell(ws,r,3,round(st["exp"],2), fmt="0.00", align="center", fill=exp_fill2)
        _cell(ws,r,4,st["n"],align="center")
        r+=1
    r+=1

    # ── Per Verdict ──
    ws.merge_cells(f"A{r}:E{r}")
    ch = ws.cell(r,1,"ANALISIS PER VERDICT")
    ch.font = Font(bold=True,color="FFFFFF"); ch.fill=HEADER_FILL
    ch.alignment = Alignment(horizontal="center")
    r+=1

    for hdr,col in [("Verdict",1),("Win Rate",2),("Expectancy (R)",3),("Sampel (n)",4)]:
        _cell(ws,r,col,hdr,bold=True,fill=PatternFill("solid",fgColor="D9E1F2"))
    r+=1

    for vrd, st in m["verdict"].items():
        exp_fill3 = GREEN_FILL if st["exp"]>=0 else RED_FILL
        _cell(ws,r,1,vrd)
        _cell(ws,r,2,f"{st['wr']*100:.1f}%", align="center")
        _cell(ws,r,3,round(st["exp"],2), fmt="0.00", align="center", fill=exp_fill3)
        _cell(ws,r,4,st["n"],align="center")
        r+=1


# ═══════════════════════════════════════════════════════════════════════════
#  TULIS SHEET EQUITY
# ═══════════════════════════════════════════════════════════════════════════

def write_equity(wb, m):
    ws = wb["Equity"]
    ws.delete_rows(1, ws.max_row + 1)

    # Hapus chart lama jika ada
    ws._charts = []

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 16

    # Header
    _cell(ws, 1, 1, "No Trade",    bold=True, fill=HEADER_FILL,
          align="center").font = Font(bold=True, color="FFFFFF")
    _cell(ws, 1, 2, "R Kumulatif", bold=True, fill=HEADER_FILL,
          align="center").font = Font(bold=True, color="FFFFFF")

    equity = m["equity"]
    for i, val in enumerate(equity, 1):
        ws.cell(i+1, 1, i)
        c = ws.cell(i+1, 2, val)
        c.number_format = "0.00"
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
        c.fill = GREEN_FILL if val >= 0 else RED_FILL

    if len(equity) >= 2:
        chart = LineChart()
        chart.title = "Equity Curve (R Kumulatif)"
        chart.style = 10
        chart.y_axis.title = "R"
        chart.x_axis.title = "Trade ke-"
        chart.width  = 20
        chart.height = 12

        data = Reference(ws, min_col=2, min_row=1, max_row=len(equity)+1)
        chart.add_data(data, titles_from_data=True)

        cats = Reference(ws, min_col=1, min_row=2, max_row=len(equity)+1)
        chart.set_categories(cats)

        chart.series[0].graphicalProperties.line.solidFill = "2F4F8F"
        chart.series[0].graphicalProperties.line.width = 20000
        ws.add_chart(chart, "D2")


# ═══════════════════════════════════════════════════════════════════════════
#  RECOMPUTE SEMUA
# ═══════════════════════════════════════════════════════════════════════════

def recompute(wb):
    rows = read_log(wb)
    m    = compute_metrics(rows)
    write_evaluasi(wb, m)
    write_equity(wb, m)
    return m


def save_wb(wb):
    try:
        wb.save(EXCEL_FILE)
    except PermissionError:
        print(f"\n[ERROR] Tidak bisa menyimpan '{EXCEL_FILE}' — file sedang dibuka di Excel.")
        print("        Tutup dulu file-nya, lalu jalankan ulang.\n")
        sys.exit(1)


# ═══════════════════════════════════════════════════════════════════════════
#  MODE A — INTERAKTIF
# ═══════════════════════════════════════════════════════════════════════════

def ask(prompt, default=None, choices=None):
    hint = ""
    if choices:
        hint = f" [{'/'.join(choices)}]"
    if default is not None:
        hint += f" (default: {default})"
    while True:
        val = input(f"  {prompt}{hint}: ").strip()
        if not val and default is not None:
            return default
        if choices and val not in choices:
            print(f"    ! Pilihan valid: {choices}")
            continue
        return val


def interactive_add(wb):
    print("\n── Catat Call Baru ──────────────────────────────")
    d = {}
    d["No"]         = next_no(wb)
    d["Tanggal"]    = datetime.now().strftime("%Y-%m-%d %H:%M")
    d["Coin"]       = ask("Coin").upper()
    d["Setup"]      = ask("Setup", choices=sorted(VALID_SETUP))
    d["Bias"]       = ask("Bias", choices=sorted(VALID_BIAS))
    d["Score"]      = ask("Score (0-100)")
    d["Confidence"] = ask("Confidence", choices=sorted(VALID_CONFIDENCE))
    d["Verdict"]    = ask("Verdict",    choices=sorted(VALID_VERDICT))
    d["Entry"]      = ask("Entry price")
    d["SL"]         = ask("SL price")
    d["TP1"]        = ask("TP1 price")
    d["TP2"]        = ask("TP2 price (kosongkan jika tidak ada)", default="")
    d["TP3"]        = ask("TP3 price (kosongkan jika tidak ada)", default="")
    d["RR_plan"]    = ask("RR plan (mis. 1:3.5)", default="")
    for p in ("P1","P2","P3","P4","P5","P6"):
        d[p] = ask(f"{p} (-1/0/1)", choices=["-1","0","1"])
    d["Diambil"]         = ask("Diambil (Y/N)", choices=["Y","N"])
    d["Hasil"]           = "Pending"
    d["Exit"]            = ""
    d["R_aktual"]        = ""
    d["Invalidasi_kena"] = ""
    d["Catatan"]         = ask("Catatan (opsional)", default="")

    errs = validate_add(d)
    if errs:
        print("\n[VALIDASI GAGAL]")
        for e in errs: print(e)
        return

    append_trade(wb, d)
    m = recompute(wb)
    save_wb(wb)
    print(f"\n✔ Trade #{d['No']} ({d['Coin']}) berhasil dicatat.")
    _print_summary(m)


def interactive_update(wb):
    print("\n── Update Hasil Trade ───────────────────────────")
    try:
        no = int(ask("Nomor trade yang ingin diupdate"))
    except ValueError:
        print("  ! Nomor harus angka."); return

    rows = read_log(wb)
    match = [r for r in rows if int(r["No"]) == no]
    if not match:
        print(f"  ! Trade #{no} tidak ditemukan."); return

    d = {}
    d["Hasil"]           = ask("Hasil", choices=sorted(VALID_HASIL))
    d["R_aktual"]        = ask("R aktual (mis. +3.0 atau -1.0)")
    d["Exit"]            = ask("Exit price")
    d["Invalidasi_kena"] = ask("Invalidasi kena? (Y/N)", choices=["Y","N"])
    d["Catatan"]         = ask("Catatan tambahan", default="")

    errs = validate_hasil(d)
    if errs:
        print("\n[VALIDASI GAGAL]")
        for e in errs: print(e)
        return

    ok = update_trade(wb, no, d)
    if ok:
        m = recompute(wb)
        save_wb(wb)
        print(f"\n✔ Trade #{no} berhasil diupdate.")
        _print_summary(m)
    else:
        print(f"  ! Gagal mengupdate trade #{no}.")


def _print_summary(m):
    print("\n── Ringkasan Cepat ──────────────────────────────")
    exp_sign = "+" if m["expectancy"] >= 0 else ""
    print(f"  Expectancy    : {exp_sign}{m['expectancy']:.2f} R")
    print(f"  Win Rate      : {m['win_rate']*100:.1f}%  ({m['wins']}W / {m['losses']}L / {m['bes']}BE)")
    print(f"  Total R       : {m['total_r']:+.2f} R")
    print(f"  Trade selesai : {m['finished']}  |  Pending: {m['pending']}")
    print("─────────────────────────────────────────────────\n")


def interactive_mode():
    wb = open_or_create_workbook()
    while True:
        print("\n╔══════════════════════════════════════════╗")
        print("║     QUANT EDGE ENGINE — Trade Journal    ║")
        print("╠══════════════════════════════════════════╣")
        print("║  1. Catat call baru                      ║")
        print("║  2. Update hasil trade                   ║")
        print("║  3. Recompute evaluasi                   ║")
        print("║  4. Lihat ringkasan                      ║")
        print("║  0. Keluar                               ║")
        print("╚══════════════════════════════════════════╝")
        pilih = input("Pilih menu: ").strip()

        if pilih == "1":
            interactive_add(wb)
        elif pilih == "2":
            interactive_update(wb)
        elif pilih == "3":
            m = recompute(wb)
            save_wb(wb)
            print("✔ Recompute selesai.")
            _print_summary(m)
        elif pilih == "4":
            rows = read_log(wb)
            m = compute_metrics(rows)
            _print_summary(m)
        elif pilih == "0":
            print("Sampai jumpa!")
            break
        else:
            print("  ! Pilihan tidak valid.")


# ═══════════════════════════════════════════════════════════════════════════
#  MODE B — SATU-BARIS CEPAT
# ═══════════════════════════════════════════════════════════════════════════

def parse_add_args(arg_str):
    """
    Format: "COIN,Setup,bias,score,confidence,verdict,entry,SL,TP1,TP2,TP3,RR,P1,P2,P3,P4,P5,P6"
    Field opsional di belakang boleh dikosongkan.
    """
    parts = [p.strip() for p in arg_str.split(",")]
    keys  = ["Coin","Setup","Bias","Score","Confidence","Verdict",
             "Entry","SL","TP1","TP2","TP3","RR_plan",
             "P1","P2","P3","P4","P5","P6"]
    d = {}
    for i, k in enumerate(keys):
        d[k] = parts[i].strip() if i < len(parts) else ""
    if d.get("Coin"):
        d["Coin"] = d["Coin"].upper()
    return d


def parse_result_args(arg_str):
    """Format: "Hasil,R_aktual,Exit,catatan" """
    parts = [p.strip() for p in arg_str.split(",", 3)]
    keys  = ["Hasil","R_aktual","Exit","Catatan"]
    d = {}
    for i, k in enumerate(keys):
        d[k] = parts[i] if i < len(parts) else ""
    return d


def cmd_add(arg_str):
    wb = open_or_create_workbook()
    d  = parse_add_args(arg_str)

    errs = validate_add(d)
    if errs:
        print("\n[VALIDASI GAGAL]")
        for e in errs: print(e)
        print("\nContoh format yang benar:")
        print('  python journal.py add "SUI,Narrative,long,78,High,EKSEKUSI,3.85,3.77,4.0,4.2,4.5,1:3.5,1,1,1,1,1,0"')
        sys.exit(1)

    d["No"]      = next_no(wb)
    d["Tanggal"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    d["Hasil"]   = "Pending"
    d.setdefault("Diambil", "")
    d.setdefault("Exit", "")
    d.setdefault("R_aktual", "")
    d.setdefault("Invalidasi_kena", "")
    d.setdefault("Catatan", "")

    append_trade(wb, d)
    m = recompute(wb)
    save_wb(wb)
    print(f"✔ Trade #{d['No']} ({d['Coin']}) dicatat — {EXCEL_FILE} diupdate.")
    _print_summary(m)


def cmd_result(no_str, arg_str):
    try:
        no = int(no_str)
    except ValueError:
        print(f"[ERROR] Nomor trade harus angka, bukan '{no_str}'")
        sys.exit(1)

    wb = open_or_create_workbook()
    d  = parse_result_args(arg_str)

    errs = validate_hasil(d)
    if errs:
        print("\n[VALIDASI GAGAL]")
        for e in errs: print(e)
        print("\nContoh format yang benar:")
        print('  python journal.py result 7 "Win,3.0,4.5,TP3 kena bersih"')
        sys.exit(1)

    ok = update_trade(wb, no, d)
    if not ok:
        rows = read_log(wb)
        nos  = [str(r["No"]) for r in rows]
        print(f"[ERROR] Trade #{no} tidak ditemukan. No yang ada: {', '.join(nos) if nos else '(kosong)'}")
        sys.exit(1)

    m = recompute(wb)
    save_wb(wb)
    print(f"✔ Trade #{no} diupdate — {EXCEL_FILE} diupdate.")
    _print_summary(m)


def cmd_stats():
    wb   = open_or_create_workbook()
    rows = read_log(wb)
    m    = compute_metrics(rows)
    print("\n── Quant Edge Engine — Stats ────────────────────")
    _print_summary(m)

    if m["pilar"]:
        print("  Pilar  │ WR      │ n   │ Ket")
        print("  ────────────────────────────────────────────")
        for p in ("P1","P2","P3","P4","P5","P6"):
            st = m["pilar"][p]
            if st["wr"] is None:
                print(f"  {p}     │ —       │ {st['n']:<3} │ belum ada data")
            else:
                wr_s = f"{st['wr']*100:.1f}%"
                sml  = " (kecil)" if st["n"] < 10 else ""
                print(f"  {p}     │ {wr_s:<7} │ {st['n']:<3} │{sml}")
    print()


def cmd_recompute():
    wb = open_or_create_workbook()
    m  = recompute(wb)
    save_wb(wb)
    print(f"✔ Recompute selesai — {EXCEL_FILE} diupdate.")
    _print_summary(m)


# ═══════════════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ═══════════════════════════════════════════════════════════════════════════

def main():
    args = sys.argv[1:]

    if not args:
        # Mode A — interaktif
        interactive_mode()
        return

    cmd = args[0].lower()

    if cmd == "add":
        if len(args) < 2:
            print('[ERROR] Butuh argumen. Contoh:')
            print('  python journal.py add "SUI,Narrative,long,78,High,EKSEKUSI,3.85,3.77,4.0,4.2,4.5,1:3.5,1,1,1,1,1,0"')
            sys.exit(1)
        cmd_add(args[1])

    elif cmd == "result":
        if len(args) < 3:
            print('[ERROR] Butuh 2 argumen. Contoh:')
            print('  python journal.py result 7 "Win,3.0,4.5,TP3 kena bersih"')
            sys.exit(1)
        cmd_result(args[1], args[2])

    elif cmd == "stats":
        cmd_stats()

    elif cmd == "recompute":
        cmd_recompute()

    else:
        print(f"[ERROR] Perintah '{cmd}' tidak dikenal.")
        print("Perintah tersedia: add, result, stats, recompute")
        print("Jalankan tanpa argumen untuk mode interaktif.")
        sys.exit(1)


if __name__ == "__main__":
    main()
